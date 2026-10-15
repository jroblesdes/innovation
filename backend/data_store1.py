"""Loads the project data files once and exposes simple lookups.

Files (in ../data):
  patient_charts.json        JSON Lines, one fictional MIMIC-IV chart per line
  transfer_scenarios.csv     transfer requests (department + requested level of care)
  bed_board_and_status.csv   current bed board snapshot
  RPRB_Rule_Set.xlsx         specialty -> unit routing rules
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DATA_DIR = Path(__file__).resolve().parent.parent / "data"

# Bed status codes used on the bed board
BED_STATUS = {
    "A": "Available",
    "Oc": "Occupied",
    "D": "Dirty",
    "B": "Blocked",
    "OOS": "Out of service",
}


@dataclass
class Bed:
    department: str
    unit: str
    room: str
    bed: str
    status: str
    patient_sex: str
    needs_isolation: str

    @property
    def bed_id(self) -> str:
        return f"{self.unit}-{self.room}{self.bed}"

    @property
    def is_available(self) -> bool:
        return self.status == "A"

    @property
    def is_occupied(self) -> bool:
        return self.status == "Oc"


@dataclass
class TransferRequest:
    transfer_id: str
    hadm_id: str
    subject_id: str
    department: str
    level_of_care: str
    requested_level_of_care: str
    careunit: str
    intime: str
    outtime: str
    curr_service: str
    provider_name: str
    services: list[str] = field(default_factory=list)


class DataStore:
    def __init__(self, data_dir: Path = DATA_DIR) -> None:
        self.data_dir = data_dir
        self.charts = self._load_charts()
        self.transfers = self._load_transfers()
        self.beds = self._load_beds()
        self.rules = self._load_rules()

    # ---------- loaders ----------
    def _load_charts(self) -> dict[str, dict[str, Any]]:
        charts: dict[str, dict[str, Any]] = {}
        with (self.data_dir / "patient_charts.json").open(encoding="utf-8") as fh:
            for line in fh:
                if line.strip():
                    chart = json.loads(line)
                    charts[str(chart["patient"]["subject_id"])] = chart
        return charts

    def _load_transfers(self) -> dict[str, TransferRequest]:
        """One request per transfer_id.

        The CSV repeats a transfer once per service change; those rows only differ
        in prev/curr service and transfer time, so they are merged here.
        """
        out: dict[str, TransferRequest] = {}
        with (self.data_dir / "transfer_scenarios.csv").open(encoding="utf-8") as fh:
            for row in csv.DictReader(fh):
                tid = row["transfer_id"]
                if tid not in out:
                    out[tid] = TransferRequest(
                        transfer_id=tid,
                        hadm_id=row["hadm_id"],
                        subject_id=row["subject_id"],
                        department=row["department"],
                        level_of_care=row["level_of_care"],
                        requested_level_of_care=row["requested_level_of_care"],
                        careunit=row["careunit"],
                        intime=row["intime"],
                        outtime=row["outtime"],
                        curr_service=row["curr_service"],
                        provider_name=row["provider_name"],
                    )
                req = out[tid]
                svc = row["curr_service"]
                if svc and svc not in req.services:
                    req.services.append(svc)
                # keep the most recent service as the current one
                req.curr_service = svc or req.curr_service
        return dict(sorted(out.items(), key=lambda kv: kv[1].intime))

    def _load_beds(self) -> list[Bed]:
        with (self.data_dir / "bed_board_and_status.csv").open(encoding="utf-8") as fh:
            return [
                Bed(
                    department=r["full_dept_name"],
                    unit=r["colloquial_name"],
                    room=r["room_number"],
                    bed=r["bed_identifier"],
                    status=r["bed_status"],
                    patient_sex=r["patient_sex"],
                    needs_isolation=r["needs_isolation"],
                )
                for r in csv.DictReader(fh)
            ]

    def _load_rules(self) -> dict[str, dict[str, list[str]]]:
        """RPRB rules: specialty -> {column -> [units]}."""
        ws = load_workbook(self.data_dir / "RPRB_Rule_Set.xlsx", read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip() for h in rows[0]]
        rules: dict[str, dict[str, list[str]]] = {}
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            specialty = str(row[0]).strip()
            rules[specialty] = {
                header[i]: [u.strip() for u in str(v).split(",") if u.strip()]
                if v
                else []
                for i, v in enumerate(row)
                if i > 0
            }
        return rules

    # ---------- lookups ----------
    def chart(self, subject_id: str) -> dict[str, Any]:
        return self.charts[str(subject_id)]

    def beds_in_room(self, unit: str, room: str) -> list[Bed]:
        return [b for b in self.beds if b.unit == unit and b.room == room]

    def bed_board_units(self) -> set[str]:
        return {b.unit for b in self.beds}


@lru_cache(maxsize=1)
def get_store() -> DataStore:
    return DataStore()
